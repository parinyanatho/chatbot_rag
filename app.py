"""
Chatbot RAG ภาษาไทยสำหรับ Streamlit
- ผู้ใช้อัปโหลดเอกสาร (PDF, DOCX, TXT, MD, CSV) เพื่อสร้างฐานความรู้เองได้
- หากไม่พบข้อมูลที่เกี่ยวข้องในเอกสาร ระบบจะตอบด้วย LLM ฟรีที่ไม่ต้องใช้ API key
"""
import io
import threading

import numpy as np
import requests
import streamlit as st

st.set_page_config(page_title="RAG Chatbot", page_icon="💬", layout="wide")

# ---------------------------------------------------------------- ค่าตั้งต้น
EMBED_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"

LLM_OPTIONS = {
    "Qwen2.5-0.5B-Instruct (รันในเซิร์ฟเวอร์, เบา)": "Qwen/Qwen2.5-0.5B-Instruct",
    "Qwen2.5-1.5B-Instruct (รันในเซิร์ฟเวอร์, ใช้ RAM มาก)": "Qwen/Qwen2.5-1.5B-Instruct",
    "Pollinations (ฟรี ไม่ใช้ API key, เรียกผ่านอินเทอร์เน็ต)": "pollinations",
}
POLLINATIONS_URL = "https://text.pollinations.ai/openai"

SYSTEM_RAG = (
    "คุณเป็นผู้ช่วยตอบคำถามภาษาไทยที่สุภาพและชัดเจน "
    "ให้ตอบโดยอ้างอิงจาก 'ข้อมูลอ้างอิง' ที่ให้มาเท่านั้น "
    "หากข้อมูลไม่เพียงพอ ให้แจ้งตามตรงว่าไม่พบในเอกสาร ห้ามแต่งข้อมูลเพิ่ม"
)
SYSTEM_GENERAL = (
    "คุณเป็นผู้ช่วยภาษาไทยที่สุภาพและชัดเจน ตอบกระชับ ตรงประเด็น "
    "หากไม่แน่ใจให้บอกข้อจำกัดของคำตอบ ห้ามแต่งข้อมูลหรือแหล่งอ้างอิง"
)


# ---------------------------------------------------------------- โมเดล
@st.cache_resource(show_spinner="กำลังโหลดโมเดล Embedding...")
def load_embedder():
    from sentence_transformers import SentenceTransformer

    return SentenceTransformer(EMBED_MODEL)


@st.cache_resource(show_spinner="กำลังโหลดโมเดลภาษา (ครั้งแรกอาจใช้เวลาหลายนาที)...")
def load_llm(model_id: str):
    from transformers import AutoModelForCausalLM, AutoTokenizer

    tok = AutoTokenizer.from_pretrained(model_id)
    model = AutoModelForCausalLM.from_pretrained(model_id, torch_dtype="auto")
    model.eval()
    return tok, model


def generate_local(model_id, messages, max_new_tokens, temperature):
    """สร้างข้อความจากโมเดลในเครื่องแบบสตรีม"""
    from transformers import TextIteratorStreamer

    tok, model = load_llm(model_id)
    prompt = tok.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
    inputs = tok([prompt], return_tensors="pt")
    streamer = TextIteratorStreamer(tok, skip_prompt=True, skip_special_tokens=True)
    kwargs = dict(
        **inputs,
        streamer=streamer,
        max_new_tokens=max_new_tokens,
        repetition_penalty=1.1,
    )
    if temperature > 0:
        kwargs.update(do_sample=True, temperature=temperature, top_p=0.9)
    else:
        kwargs.update(do_sample=False)
    threading.Thread(target=model.generate, kwargs=kwargs, daemon=True).start()
    for piece in streamer:
        yield piece


def generate_pollinations(messages):
    """เรียกบริการฟรีที่ไม่ต้องใช้ API key (ขึ้นกับความพร้อมของผู้ให้บริการ)"""
    r = requests.post(
        POLLINATIONS_URL,
        json={"model": "openai", "messages": messages},
        timeout=60,
    )
    r.raise_for_status()
    yield r.json()["choices"][0]["message"]["content"]


def generate(backend_key, messages, max_new_tokens, temperature):
    """เลือก backend และสำรองด้วยโมเดลในเครื่องหากบริการภายนอกใช้งานไม่ได้"""
    if backend_key == "pollinations":
        try:
            yield from generate_pollinations(messages)
            return
        except Exception:
            yield "_(บริการภายนอกไม่พร้อมใช้งาน จึงสลับไปใช้โมเดลในเครื่อง)_\n\n"
            backend_key = "Qwen/Qwen2.5-0.5B-Instruct"
    yield from generate_local(backend_key, messages, max_new_tokens, temperature)


# ---------------------------------------------------------------- เอกสาร / RAG
def read_file(name: str, data: bytes) -> str:
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == "docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    # txt, md, csv
    for enc in ("utf-8", "utf-8-sig", "cp874", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return ""


def chunk_text(text: str, size: int, overlap: int):
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    chunks, start, n = [], 0, len(text)
    while start < n:
        end = min(start + size, n)
        if end < n:  # พยายามตัดที่ขึ้นบรรทัดใหม่หรือช่องว่างใกล้ท้าย chunk
            cut = max(text.rfind("\n", start, end), text.rfind(" ", start, end))
            if cut > start + size * 0.6:
                end = cut
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return chunks


def init_state():
    st.session_state.setdefault("messages", [])
    st.session_state.setdefault("chunks", [])  # [{"text":..., "source":...}]
    st.session_state.setdefault("emb", None)
    st.session_state.setdefault("files", set())


def add_documents(files, size, overlap):
    embedder = load_embedder()
    added = 0
    for f in files:
        if f.name in st.session_state.files:
            continue
        text = read_file(f.name, f.getvalue())
        pieces = chunk_text(text, size, overlap)
        if not pieces:
            st.warning(f"ไม่พบข้อความที่อ่านได้ในไฟล์ {f.name} (ไฟล์สแกนภาพต้องทำ OCR ก่อน)")
            continue
        emb = embedder.encode(pieces, normalize_embeddings=True, show_progress_bar=False)
        st.session_state.chunks += [{"text": p, "source": f.name} for p in pieces]
        prev = st.session_state.emb
        st.session_state.emb = emb if prev is None else np.vstack([prev, emb])
        st.session_state.files.add(f.name)
        added += len(pieces)
    return added


def retrieve(query, top_k):
    if st.session_state.emb is None:
        return []
    q = load_embedder().encode([query], normalize_embeddings=True)[0]
    scores = st.session_state.emb @ q
    idx = np.argsort(-scores)[:top_k]
    return [(float(scores[i]), st.session_state.chunks[i]) for i in idx]


# ---------------------------------------------------------------- UI
init_state()

with st.sidebar:
    st.header("⚙️ การตั้งค่า")
    backend_label = st.selectbox("โมเดลภาษา (LLM)", list(LLM_OPTIONS.keys()))
    backend_key = LLM_OPTIONS[backend_label]
    max_tokens = st.slider("ความยาวคำตอบสูงสุด (tokens)", 128, 1024, 400, 32)
    temperature = st.slider("Temperature", 0.0, 1.0, 0.3, 0.1)

    st.divider()
    st.subheader("📚 ฐานความรู้ (RAG)")
    uploaded = st.file_uploader(
        "อัปโหลดเอกสาร",
        type=["pdf", "docx", "txt", "md", "csv"],
        accept_multiple_files=True,
    )
    chunk_size = st.slider("ขนาด chunk (ตัวอักษร)", 300, 2000, 800, 50)
    chunk_overlap = st.slider("ส่วนซ้อนทับ (ตัวอักษร)", 0, 400, 100, 10)
    top_k = st.slider("จำนวนส่วนที่ดึงมาอ้างอิง (Top-K)", 1, 8, 3)
    threshold = st.slider(
        "เกณฑ์ความเกี่ยวข้องขั้นต่ำ", 0.0, 1.0, 0.35, 0.05,
        help="หากคะแนนสูงสุดต่ำกว่าเกณฑ์นี้ ระบบจะตอบจากความรู้ทั่วไปของ LLM",
    )

    if st.button("เพิ่มเอกสารเข้าฐานความรู้", type="primary", disabled=not uploaded):
        with st.spinner("กำลังประมวลผลเอกสาร..."):
            n = add_documents(uploaded, chunk_size, chunk_overlap)
        st.success(f"เพิ่มข้อมูลใหม่ {n} ส่วน") if n else st.info("ไม่มีเอกสารใหม่ให้เพิ่ม")

    if st.session_state.files:
        st.caption(f"เอกสารในระบบ {len(st.session_state.files)} ไฟล์ / {len(st.session_state.chunks)} ส่วน")
        for name in sorted(st.session_state.files):
            st.caption(f"• {name}")
    if st.button("ล้างฐานความรู้"):
        st.session_state.chunks, st.session_state.emb = [], None
        st.session_state.files = set()
        st.rerun()
    if st.button("ล้างประวัติการสนทนา"):
        st.session_state.messages = []
        st.rerun()

st.title("💬 Chatbot RAG")
st.caption("ถามจากเอกสารที่คุณอัปโหลด หากไม่พบข้อมูลจะตอบด้วย LLM ฟรีที่ไม่ต้องใช้ API key")

for m in st.session_state.messages:
    with st.chat_message(m["role"]):
        st.markdown(m["content"])
        if m.get("sources"):
            with st.expander("แหล่งข้อมูลอ้างอิง"):
                for s in m["sources"]:
                    st.markdown(f"**{s['source']}** (คะแนน {s['score']:.2f})\n\n> {s['text'][:400]}...")

if question := st.chat_input("พิมพ์คำถามของคุณ"):
    st.session_state.messages.append({"role": "user", "content": question})
    with st.chat_message("user"):
        st.markdown(question)

    hits = retrieve(question, top_k)
    use_rag = bool(hits) and hits[0][0] >= threshold
    history = [
        {"role": m["role"], "content": m["content"]}
        for m in st.session_state.messages[:-1][-6:]
    ]

    if use_rag:
        context = "\n\n".join(
            f"[{i}] (ไฟล์: {h['source']})\n{h['text']}" for i, (_, h) in enumerate(hits, 1)
        )
        user_msg = f"ข้อมูลอ้างอิง:\n{context}\n\nคำถาม: {question}"
        messages = [{"role": "system", "content": SYSTEM_RAG}] + history + [{"role": "user", "content": user_msg}]
    else:
        messages = [{"role": "system", "content": SYSTEM_GENERAL}] + history + [{"role": "user", "content": question}]

    with st.chat_message("assistant"):
        if not use_rag:
            st.info("ไม่พบข้อมูลที่เกี่ยวข้องในเอกสาร จึงตอบจากความรู้ทั่วไปของโมเดล ควรตรวจสอบความถูกต้องอีกครั้ง")
        try:
            answer = st.write_stream(generate(backend_key, messages, max_tokens, temperature))
        except Exception as e:
            answer = f"เกิดข้อผิดพลาดในการสร้างคำตอบ: {e}"
            st.error(answer)
        sources = [{"score": s, **h} for s, h in hits] if use_rag else []
        if sources:
            with st.expander("แหล่งข้อมูลอ้างอิง"):
                for s in sources:
                    st.markdown(f"**{s['source']}** (คะแนน {s['score']:.2f})\n\n> {s['text'][:400]}...")

    st.session_state.messages.append({"role": "assistant", "content": answer, "sources": sources})
